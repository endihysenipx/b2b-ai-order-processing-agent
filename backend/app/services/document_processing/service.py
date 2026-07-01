import csv
from pathlib import Path

import openpyxl
import pdfplumber
from docx import Document


def detect_file_type(file_path: str) -> str:
    return Path(file_path).suffix.lower().lstrip(".")


def extract_pdf_text(file_path: str) -> str:
    with pdfplumber.open(file_path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def extract_word_text(file_path: str) -> str:
    document = Document(file_path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def extract_excel_data(file_path: str) -> list[list[str]]:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    sheet = workbook.active
    return [[str(cell) if cell is not None else "" for cell in row] for row in sheet.iter_rows(values_only=True)]


def extract_csv_data(file_path: str) -> list[list[str]]:
    with open(file_path, newline="", encoding="utf-8") as handle:
        return list(csv.reader(handle))


def extract_image_text(file_path: str) -> str:
    return f"OCR placeholder for image file {file_path}"
